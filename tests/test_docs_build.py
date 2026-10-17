"""Тесты сборки программной документации (``docs/build_docx.py``).

Проверяется, что преобразователь Markdown → DOCX разбирает заголовки, таблицы,
блоки кода и выделение текста, а также что настоящие разделы документации
собираются в один файл без ошибок.
"""

from __future__ import annotations

import importlib.util
from pathlib import Path
from types import ModuleType

import pytest

pytest.importorskip("docx")

from docx import Document  # noqa: E402

PROJECT_ROOT = Path(__file__).resolve().parents[1]
BUILDER_PATH = PROJECT_ROOT / "docs" / "build_docx.py"

DOC_SOURCES = ("theory.md", "user_guide.md", "api.md")

SAMPLE_MARKDOWN = """# Раздел

Абзац с **жирным**, `кодом` и *курсивом*.

| Ключ | Значение |
| --- | --- |
| a | 1 |
| b | 2 |

- пункт один
- пункт два

```text
A + B = [3, 8]
```

1. первый
2. второй

> примечание
"""


@pytest.fixture(scope="module")
def builder() -> ModuleType:
    """Загрузить скрипт сборки как модуль."""
    spec = importlib.util.spec_from_file_location("docs_build_docx", BUILDER_PATH)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


@pytest.fixture()
def sample_markdown(tmp_path: Path) -> Path:
    path = tmp_path / "sample.md"
    path.write_text(SAMPLE_MARKDOWN, encoding="utf-8")
    return path


def test_build_from_sample(builder: ModuleType, sample_markdown: Path, tmp_path: Path) -> None:
    output = builder.build_document([sample_markdown], tmp_path / "sample.docx")
    assert output.exists()

    document = Document(str(output))
    styles = [paragraph.style.name for paragraph in document.paragraphs]

    assert "Heading 1" in styles
    assert "Код" in styles
    assert "Цитата" in styles
    assert [paragraph.text for paragraph in document.paragraphs if paragraph.style.name == "Heading 1"] == [
        "Содержание",
        "Раздел",
    ]

    # Первая таблица — таблица из Markdown с жирной шапкой.
    table = document.tables[0]
    assert [cell.text for cell in table.rows[0].cells] == ["Ключ", "Значение"]
    assert [cell.text for cell in table.rows[1].cells] == ["a", "1"]
    assert all(run.bold for run in table.rows[0].cells[0].paragraphs[0].runs)

    # Строка кода выводится как есть, моноширинным стилем.
    code = next(p for p in document.paragraphs if p.style.name == "Код")
    assert code.text == 'A + B = [3, 8]'
    assert document.styles["Код"].font.name == "Consolas"


def test_markdown_is_not_left_in_text(builder: ModuleType, sample_markdown: Path, tmp_path: Path) -> None:
    output = builder.build_document([sample_markdown], tmp_path / "sample.docx")
    document = Document(str(output))

    for paragraph in document.paragraphs:
        if paragraph.style.name == "Код":
            continue
        assert "**" not in paragraph.text
        assert "```" not in paragraph.text
        assert not paragraph.text.strip().startswith("|")

    bullet = next(p for p in document.paragraphs if p.text.startswith("• пункт один"))
    assert bullet.runs[1].bold is None
    bold = next(p for p in document.paragraphs if p.text.startswith("Абзац с"))
    assert any(run.bold for run in bold.runs)


def test_real_documentation_builds(builder: ModuleType, tmp_path: Path) -> None:
    sources = [PROJECT_ROOT / "docs" / name for name in DOC_SOURCES]
    for source in sources:
        assert source.exists(), f"нет файла документации: {source}"

    output = builder.build_document(sources, tmp_path / "documentation.docx")
    document = Document(str(output))

    top_level = [p.text for p in document.paragraphs if p.style.name == "Heading 1"]
    assert top_level[0] == "Содержание"
    assert len(top_level) == 4  # содержание + три раздела
    assert len(document.tables) >= 25

    # Содержание — обычный текст: разделы и подразделы перечислены в документе.
    contents = [p.text for p in document.paragraphs if p.text.startswith("Теоретический раздел")]
    assert contents and contents[0] == "Теоретический раздел"
    assert any(p.text.startswith("1. Нечеткая величина") for p in document.paragraphs)
    assert not any("F9" in p.text for p in document.paragraphs)
