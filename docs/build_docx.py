"""Сборка программной документации: Markdown → DOCX.

Скрипт превращает три раздела документации (``theory.md``, ``user_guide.md``,
``api.md``) в один файл Word: титульного листа нет, документ начинается с
заголовка, оглавления и далее идёт по разделам, каждый с новой страницы.

Что поддерживается из Markdown:

* заголовки ``#`` … ``####`` (стили Heading 1…4);
* абзацы с выделением ``**жирный**``, ``*курсив*`` и моноширинным ``кодом``;
* маркированные (``-``) и нумерованные (``1.``) списки с вложенностью;
* блоки кода ``` ``` ``` — моноширинным шрифтом без отбивки между строками;
* таблицы в формате «pipe»;
* цитаты (``>``) — курсивом с отступом.

Содержание по умолчанию — обычный список заголовков первого и второго уровня:
он виден сразу, ничего обновлять не нужно. Ключ ``--toc field`` вставляет вместо
него поле TOC Word (тогда номера страниц подставляет Word при обновлении поля),
``--toc none`` — собирает документ без содержания.

Запуск из корня проекта::

    .venv\\Scripts\\python.exe docs\\build_docx.py
    .venv\\Scripts\\python.exe docs\\build_docx.py --output docs/документация.docx
    .venv\\Scripts\\python.exe docs\\build_docx.py --toc field
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path
from typing import Iterable, Iterator, List, Sequence

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DOCS_DIR = PROJECT_ROOT / "docs"

#: Разделы документации в порядке следования в итоговом файле.
DEFAULT_SOURCES: Sequence[str] = ("theory.md", "user_guide.md", "api.md")

#: Имя итогового файла по умолчанию.
DEFAULT_OUTPUT = "Неклассические_логики_документация.docx"

DOC_TITLE = (
    "Неклассические логики — нечеткие величины "
    "(вариант 1: равномерное распределение)"
)
DOC_SUBTITLE = "Программная документация"

BODY_FONT = "Times New Roman"
CODE_FONT = "Consolas"
BODY_SIZE = 11.0
CODE_SIZE = 9.0

BULLET_LEVELS = ("•", "◦", "▪")
INLINE_PATTERN = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*`]+?\*)")
HEADING_PATTERN = re.compile(r"^(#{1,4})\s+(.*)$")
BULLET_PATTERN = re.compile(r"^(\s*)[-*]\s+(.*)$")
NUMBER_PATTERN = re.compile(r"^(\s*)(\d+)[.)]\s+(.*)$")
TABLE_SEPARATOR = re.compile(r"^\|?[\s:|-]*\|[\s:|-]*$")
RULE_PATTERN = re.compile(r"^\s*-{3,}\s*$")


# ----------------------------------------------------------------------
# Настройка стилей
# ----------------------------------------------------------------------
def _force_fonts(style_element, name: str) -> None:
    """Прописать шрифт для латиницы, кириллицы и восточных языков."""
    rpr = style_element.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        style_element.append(rpr)
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attribute in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attribute), name)


def setup_document() -> Document:
    """Создать документ с настройками страницы и стилями."""
    document = Document()

    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)

    normal = document.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(BODY_SIZE)
    _force_fonts(normal.element, BODY_FONT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for name, size in (
        ("Heading 1", 16.0),
        ("Heading 2", 13.5),
        ("Heading 3", 12.0),
        ("Heading 4", 11.0),
    ):
        style = document.styles[name]
        style.font.name = BODY_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
        _force_fonts(style.element, BODY_FONT)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style.paragraph_format.keep_with_next = True

    code = document.styles.add_style("Код", WD_STYLE_TYPE.PARAGRAPH)
    code.base_style = document.styles["Normal"]
    code.font.name = CODE_FONT
    code.font.size = Pt(CODE_SIZE)
    _force_fonts(code.element, CODE_FONT)
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.line_spacing = 1.0
    code.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    code.paragraph_format.left_indent = Cm(0.5)

    quote = document.styles.add_style("Цитата", WD_STYLE_TYPE.PARAGRAPH)
    quote.base_style = document.styles["Normal"]
    quote.font.italic = True
    quote.paragraph_format.left_indent = Cm(1.0)

    return document


def add_page_numbers(document: Document) -> None:
    """Номер страницы в нижнем колонтитуле."""
    paragraph = document.sections[0].footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instruction)
    run._r.append(end)


def add_toc_field(document: Document) -> None:
    """Поле оглавления Word (номера страниц подставляет сам Word)."""
    paragraph = document.add_paragraph()
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = 'TOC \\o "1-2" \\h \\z \\u'
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Оглавление появится после обновления поля (F9)."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instruction, separate, placeholder, end):
        run._r.append(element)


def allow_field_update(document: Document) -> None:
    """Попросить Word обновить поля (оглавление, номера страниц) при открытии."""
    settings = document.settings.element
    if settings.find(qn("w:updateFields")) is None:
        element = OxmlElement("w:updateFields")
        element.set(qn("w:val"), "true")
        settings.append(element)


def collect_headings(sources: Iterable[Path]) -> List[tuple]:
    """Собрать заголовки первого и второго уровня из файлов документации."""
    entries: List[tuple] = []
    for source in sources:
        for line in source.read_text(encoding="utf-8").splitlines():
            match = HEADING_PATTERN.match(line)
            if match and len(match.group(1)) <= 2:
                entries.append((len(match.group(1)), strip_markdown(match.group(2))))
    return entries


def add_contents(document: Document, entries: Sequence[tuple]) -> None:
    """Обычный (неполевой) список содержания: виден сразу, без обновления."""
    for level, text in entries:
        paragraph = document.add_paragraph()
        formatting = paragraph.paragraph_format
        formatting.alignment = WD_ALIGN_PARAGRAPH.LEFT
        formatting.left_indent = Cm(0.7 * (level - 1))
        formatting.first_line_indent = Cm(0)
        formatting.space_after = Pt(2)
        formatting.space_before = Pt(6) if level == 1 else Pt(0)
        run = paragraph.add_run(text)
        run.bold = level == 1


# ----------------------------------------------------------------------
# Разбор строк
# ----------------------------------------------------------------------
def add_inline(paragraph, text: str) -> None:
    """Добавить текст с поддержкой ``**жирного**``, ``*курсива*`` и ``кода``."""
    for part in INLINE_PATTERN.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.font.name = CODE_FONT
            run.font.size = Pt(CODE_SIZE + 0.5)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def add_code_block(document: Document, lines: Sequence[str]) -> None:
    for line in lines:
        paragraph = document.add_paragraph(style="Код")
        # Внутри блока кода разметка не разбирается: текст выводится как есть.
        paragraph.add_run(line.replace("\t", "    "))


def strip_markdown(text: str) -> str:
    """Убрать из строки разметку выделения (для заголовков)."""
    return text.replace("**", "").replace("`", "").replace("*", "").strip()


def add_table(document: Document, rows: Sequence[Sequence[str]]) -> None:
    if not rows:
        return
    width = max(len(row) for row in rows)
    table = document.add_table(rows=0, cols=width)
    table.style = "Table Grid"
    table.autofit = True
    for index, row in enumerate(rows):
        cells = table.add_row().cells
        for position in range(width):
            text = row[position] if position < len(row) else ""
            paragraph = cells[position].paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_inline(paragraph, text)
            if index == 0:
                for run in paragraph.runs:
                    run.bold = True
    document.add_paragraph()


def split_table_row(line: str) -> List[str]:
    stripped = line.strip()
    if stripped.startswith("|"):
        stripped = stripped[1:]
    if stripped.endswith("|"):
        stripped = stripped[:-1]
    return [cell.strip() for cell in stripped.split("|")]


def _is_table_start(lines: Sequence[str], index: int) -> bool:
    if index + 1 >= len(lines):
        return False
    current = lines[index].strip()
    following = lines[index + 1].strip()
    return (
        current.startswith("|")
        and following.startswith("|")
        and TABLE_SEPARATOR.match(following) is not None
        and "-" in following
    )


def convert_markdown(document: Document, text: str) -> None:
    """Перенести содержимое Markdown-файла в документ."""
    lines = text.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if not stripped:
            index += 1
            continue

        if stripped.startswith("```"):
            index += 1
            block: List[str] = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                block.append(lines[index].rstrip())
                index += 1
            index += 1  # закрывающая линия
            add_code_block(document, block)
            continue

        if _is_table_start(lines, index):
            rows: List[List[str]] = []
            while index < len(lines) and "|" in lines[index]:
                if TABLE_SEPARATOR.match(lines[index].strip()) and "-" in lines[index]:
                    index += 1
                    continue
                rows.append(split_table_row(lines[index]))
                index += 1
            add_table(document, rows)
            continue

        heading = HEADING_PATTERN.match(line)
        if heading:
            level = len(heading.group(1))
            document.add_heading(strip_markdown(heading.group(2)), level=level)
            index += 1
            continue

        if RULE_PATTERN.match(line):
            index += 1
            continue

        if stripped.startswith(">"):
            quote_lines: List[str] = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                quote_lines.append(lines[index].strip().lstrip(">").strip())
                index += 1
            paragraph = document.add_paragraph(style="Цитата")
            add_inline(paragraph, " ".join(part for part in quote_lines if part))
            continue

        bullet = BULLET_PATTERN.match(line)
        if bullet:
            depth = min(len(bullet.group(1)) // 2, len(BULLET_LEVELS) - 1)
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.6 + 0.6 * depth)
            paragraph.paragraph_format.first_line_indent = Cm(-0.6)
            paragraph.paragraph_format.space_after = Pt(3)
            paragraph.add_run(f"{BULLET_LEVELS[depth]} ")
            add_inline(paragraph, bullet.group(2).strip())
            index += 1
            continue

        numbered = NUMBER_PATTERN.match(line)
        if numbered:
            depth = min(len(numbered.group(1)) // 2, 2)
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.8 + 0.7 * depth)
            paragraph.paragraph_format.first_line_indent = Cm(-0.8)
            paragraph.paragraph_format.space_after = Pt(3)
            paragraph.add_run(f"{numbered.group(2)}. ")
            add_inline(paragraph, numbered.group(3).strip())
            index += 1
            continue

        # Обычный абзац: склеиваем строки до пустой или до начала нового блока.
        paragraph_lines: List[str] = []
        while index < len(lines):
            candidate = lines[index]
            candidate_stripped = candidate.strip()
            if not candidate_stripped:
                break
            if (
                candidate_stripped.startswith("```")
                or candidate_stripped.startswith(">")
                or HEADING_PATTERN.match(candidate)
                or BULLET_PATTERN.match(candidate)
                or NUMBER_PATTERN.match(candidate)
                or RULE_PATTERN.match(candidate)
                or _is_table_start(lines, index)
            ):
                break
            paragraph_lines.append(candidate_stripped)
            index += 1
        if paragraph_lines:
            paragraph = document.add_paragraph()
            add_inline(paragraph, " ".join(paragraph_lines))


# ----------------------------------------------------------------------
# Сборка
# ----------------------------------------------------------------------
def build_document(
    sources: Iterable[Path],
    output: Path,
    *,
    title: str = DOC_TITLE,
    subtitle: str = DOC_SUBTITLE,
    toc: str = "text",
) -> Path:
    """Собрать итоговый DOCX из списка Markdown-файлов.

    ``toc`` задаёт вид содержания: ``"text"`` — обычный список заголовков
    (виден сразу), ``"field"`` — поле TOC Word (номера страниц, нужно F9),
    ``"none"`` — без содержания.
    """
    sources = list(sources)
    document = setup_document()
    add_page_numbers(document)

    if title:
        heading = document.add_paragraph(style="Title")
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        heading.add_run(title)
    if subtitle:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(subtitle)
        run.bold = True

    if toc != "none":
        document.add_heading("Содержание", level=1)
        if toc == "field":
            allow_field_update(document)
            add_toc_field(document)
        else:
            add_contents(document, collect_headings(sources))

    for source in sources:
        document.add_page_break()
        convert_markdown(document, source.read_text(encoding="utf-8"))

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output))
    return output


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        prog="build_docx.py",
        description="Сборка документации (Markdown → DOCX)",
    )
    parser.add_argument(
        "--output",
        default=str(DOCS_DIR / DEFAULT_OUTPUT),
        help="путь к итоговому файлу .docx",
    )
    parser.add_argument(
        "--sources",
        nargs="+",
        default=list(DEFAULT_SOURCES),
        help="порядок разделов (файлы относительно docs/)",
    )
    parser.add_argument(
        "--no-title",
        action="store_true",
        help="не добавлять заголовок и содержание",
    )
    parser.add_argument(
        "--toc",
        choices=("text", "field", "none"),
        default="text",
        help=(
            "вид содержания: text — список заголовков (по умолчанию), "
            "field — поле TOC Word с номерами страниц, none — без содержания"
        ),
    )
    args = parser.parse_args(argv)

    sources = [DOCS_DIR / name for name in args.sources]
    missing = [str(path) for path in sources if not path.exists()]
    if missing:
        print(f"Не найдены файлы: {', '.join(missing)}", file=sys.stderr)
        return 2

    output = build_document(
        sources,
        Path(args.output),
        title="" if args.no_title else DOC_TITLE,
        subtitle="" if args.no_title else DOC_SUBTITLE,
        toc="none" if args.no_title else args.toc,
    )
    document = Document(str(output))
    paragraphs = len(document.paragraphs)
    tables = len(document.tables)
    size = output.stat().st_size
    print(f"Готово: {output}")
    print(f"Разделов: {len(sources)}, абзацев: {paragraphs}, таблиц: {tables}, размер: {size} байт")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
